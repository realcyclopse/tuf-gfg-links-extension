import json
import re
import sys
from pathlib import Path
import openpyxl

def slugify(text):
    if not text:
        return ""
    text = text.lower().strip()
    text = re.sub(r'[^a-z0-9]+', '-', text)
    return text.strip('-')

def main():
    script_dir = Path(__file__).resolve().parent
    ext_dir = script_dir.parent
    
    # Locate Excel file
    if len(sys.argv) > 1:
        excel_path = Path(sys.argv[1])
    else:
        # Check in extension dir, then parent dir
        excel_path = ext_dir / "TUF_Questions_Platform_Links_Verified.xlsx"
        if not excel_path.exists():
            excel_path = ext_dir.parent / "TUF_Questions_Platform_Links_Verified.xlsx"
        if not excel_path.exists():
            excel_path = ext_dir / "TUF_Questions_Platform_Links.xlsx"
        if not excel_path.exists():
            excel_path = ext_dir.parent / "TUF_Questions_Platform_Links.xlsx"

    if not excel_path.exists():
        raise FileNotFoundError(f"Could not locate verified Excel file at {excel_path}")

    # Also check if tuf_verified_problems.json exists for extra metadata (slug)
    meta_path = ext_dir / "tuf_verified_problems.json"
    if not meta_path.exists():
        meta_path = ext_dir.parent / "tuf_verified_problems.json"

    meta_by_name = {}
    if meta_path.exists():
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                meta_list = json.load(f)
                for item in meta_list:
                    q_norm = re.sub(r'[^a-z0-9]', '', item.get("question", "").lower())
                    meta_by_name[q_norm] = item
        except Exception as e:
            print(f"Warning loading metadata JSON: {e}")

    output_path = ext_dir / "data" / "problems.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    print(f"Reading verified Excel workbook from {excel_path}...")
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    sheet = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active

    problems = []
    
    for row_idx in range(2, sheet.max_row + 1):
        q_val = sheet.cell(row=row_idx, column=1).value
        if not q_val:
            continue
        question_title = str(q_val).strip()

        # Column 3 is GeeksforGeeks
        gfg_val = sheet.cell(row=row_idx, column=3).value
        gfg_url = None
        if gfg_val and str(gfg_val).strip() and str(gfg_val).strip().lower() != "not mapped":
            url_str = str(gfg_val).strip()
            if "geeksforgeeks.org" in url_str and url_str.startswith("https://"):
                gfg_url = url_str

        # Match slug from metadata if available, otherwise slugify
        q_norm = re.sub(r'[^a-z0-9]', '', question_title.lower())
        meta = meta_by_name.get(q_norm, {})
        slug = meta.get("tuf_slug") or slugify(question_title)

        problems.append({
            "question": question_title,
            "slug": slug,
            "gfg": gfg_url
        })

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(problems, f, indent=2, ensure_ascii=False)

    gfg_count = sum(1 for p in problems if p["gfg"])

    print(f"Successfully generated {output_path}")
    print(f"Total problems: {len(problems)}")
    print(f"Verified GeeksforGeeks links: {gfg_count}")

if __name__ == "__main__":
    main()
