import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

TUF_LIVE_URL = "https://takeuforward.org/prep-hub/strivers-a2z-dsa-sheet"
DATA_FILE = Path(__file__).parent / "tuf_verified_problems.json"

def main():
    if len(sys.argv) > 1:
        input_path = Path(sys.argv[1])
    elif Path("TUF_Questions_Platform_Links.xlsx").exists():
        input_path = Path("TUF_Questions_Platform_Links.xlsx")
    else:
        input_path = Path("TUF_Questions_Platform_Links_Verified.xlsx")

    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else input_path

    if not DATA_FILE.exists():
        raise FileNotFoundError(f"Verified dataset file {DATA_FILE} not found.")

    with open(DATA_FILE, "r", encoding="utf-8") as f:
        problems = json.load(f)

    print(f"Loaded {len(problems)} verified problems matching live TakeUforward practice sheet.")

    if input_path.exists():
        print(f"Loading existing template/workbook from {input_path}...")
        wb = openpyxl.load_workbook(str(input_path))
        if "Sheet1" in wb.sheetnames:
            sheet = wb["Sheet1"]
        elif "Questions" in wb.sheetnames:
            sheet = wb["Questions"]
        else:
            sheet = wb.active
    else:
        print(f"Input workbook {input_path} not found; creating new workbook...")
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = "Sheet1"

    # Remove existing tables in the sheet
    sheet.tables.clear()

    # Clear existing rows beyond header
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row)

    headers = [
        "question", "leetcode", "gfg", "codechef", "codeforces",
        "mapping_status", "source"
    ]

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    data_align = Alignment(vertical="top", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

    for r_idx, item in enumerate(problems, start=2):
        row_vals = [
            item["question"],
            item["leetcode"] if item.get("leetcode") else None,
            item["gfg"] if item.get("gfg") else None,
            item.get("codechef", "Not mapped"),
            item.get("codeforces", "Not mapped"),
            item.get("mapping_status", "No verified LC/GFG match"),
            item.get("source", TUF_LIVE_URL)
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = sheet.cell(row=r_idx, column=c_idx, value=val)
            if c_idx in (4, 5, 6):
                cell.alignment = center_align
            else:
                cell.alignment = data_align

    col_widths = {
        "A": 42.0,
        "B": 48.0,
        "C": 52.0,
        "D": 18.0,
        "E": 18.0,
        "F": 30.0,
        "G": 72.0,
    }
    for col_letter, width in col_widths.items():
        sheet.column_dimensions[col_letter].width = width

    sheet.row_dimensions[1].height = 26.0
    sheet.freeze_panes = "A2"

    table_ref = f"A1:G{len(problems) + 1}"
    try:
        table = Table(displayName="TUFPlatformLinks", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    except Exception as e:
        print(f"Warning: could not add table: {e}")

    try:
        wb.save(str(output_path))
        print(f"Wrote {len(problems)} verified problems to {output_path}")
    except PermissionError:
        alt_path = output_path.with_name(f"{output_path.stem}_Verified.xlsx")
        wb.save(str(alt_path))
        print(f"Notice: {output_path} is currently locked by Excel.")
        print(f"Wrote {len(problems)} verified problems to {alt_path} instead.")

if __name__ == "__main__":
    main()
